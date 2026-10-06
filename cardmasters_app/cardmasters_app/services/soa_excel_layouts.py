"""Designed Excel counterpart to the reference Sales Order statement."""

from io import BytesIO
import math

import frappe

from cardmasters_app.cardmasters_app.services.soa_exports import sheet_title

HEADERS = ["CUSTOMER", "INVOICE DATE", "SALES ORDER #", "Delivery Receipt#", "CSI Invoice#",
           "ITEM DESCRIPTION", "PO #", "AGED", "Remarks", "SUM of AMOUNT", "SUM of PAYMENT", "SUM of BALANCE"]
ROW_KEYS = ["customer_name", "invoice_date", "sales_order", "delivery_receipt", "csi_invoice",
            "item_description", "po_no", "aged", "remarks", "amount", "payment", "balance"]


def designed_workbook(statements):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.page import PageMargins
    from openpyxl.utils import get_column_letter
    from frappe.utils.xlsxutils import ILLEGAL_CHARACTERS_RE

    workbook = Workbook()
    workbook.remove(workbook.active)
    used = set()
    thin = Side(style="thin", color="666666")
    white = Side(style="hair", color="FFFFFF")
    grey = PatternFill("solid", fgColor="DFE4EA")
    blue = PatternFill("solid", fgColor="8396B5")
    light = PatternFill("solid", fgColor="F4F6F8")
    for statement in statements:
        sheet = workbook.create_sheet(sheet_title(statement.customer.customer_name, used))
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "F10"
        widths = [25, 18, 18, 16, 15, 33, 16, 12, 11, 15, 12, 12]
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        for row in range(1, 10):
            sheet.row_dimensions[row].height = 16
        sheet.row_dimensions[8].height = 5
        sheet.row_dimensions[9].height = 30
        settings = statement.settings

        def text(row, column, value, *, size=10, bold=False, italic=False, color="000000", align="left"):
            cell = sheet.cell(row, column)
            cell.value = ILLEGAL_CHARACTERS_RE.sub("", str(value or ""))
            cell.data_type = "s"
            cell.font = Font(name="Arial", size=size, bold=bold, italic=italic, color=color)
            cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
            return cell

        def merged(row, start, end, value, **style):
            sheet.merge_cells(start_row=row, start_column=start, end_row=row, end_column=end)
            return text(row, start, value, **style)

        logo = Image(frappe.get_app_path("cardmasters_app", "public", "images", "soa", "reference-logo.png"))
        logo.width, logo.height = 175, 98
        sheet.add_image(logo, "A1")
        sheet.merge_cells("D2:H3")
        title = text(2, 4, "STATEMENT OF ACCOUNT", size=23, bold=True, align="center")
        title.font = Font(name="Arial", size=23, bold=True, underline="single")
        chip = merged(2, 9, 12, "CUSTOMER     " + statement.customer.customer_name, size=9, color="FFFFFF")
        chip.fill = PatternFill("solid", fgColor="3D4142")
        merged(4, 9, 12, "Phone: " + (settings.get("contact_phone") or ""), size=9, bold=True)
        merged(5, 9, 12, "E-mail: " + (settings.get("contact_email") or ""), size=8, bold=True)
        merged(6, 1, 2, "Statement No:   " + str(statement.number), size=9, bold=True)
        merged(7, 1, 2, "Date:   " + statement.date.strftime("%-d-%b-%y"), size=9, bold=True)
        merged(7, 3, 12, statement.customer.customer_name, size=17, bold=True, align="center")
        for column, label in enumerate(HEADERS, 1):
            cell = text(9, column, label, size=9, bold=True, italic=column < 10,
                        color="FFFFFF" if column >= 10 else "000000", align="center")
            cell.fill = blue if column >= 10 else grey
            cell.border = Border(top=thin, bottom=thin)
        start = 10
        for offset, row in enumerate(statement.rows):
            row_number = start + offset
            sheet.row_dimensions[row_number].height = max(20, 15 * math.ceil(len(row.item_description) / 38))
            for column, key in enumerate(ROW_KEYS, 1):
                value = row[key]
                if column == 1 and offset:
                    value = ""
                if column in (3, 4, 5) and offset and not row.first_in_order and value == statement.rows[offset - 1][key]:
                    value = ""
                if column in (2, 10, 11, 12):
                    cell = sheet.cell(row_number, column, value)
                    cell.font = Font(name="Arial", size=9)
                    cell.alignment = Alignment(horizontal="right" if column >= 10 else "center", vertical="top")
                    cell.number_format = "#,##0.00" if column >= 10 else "mmmm d, yyyy"
                else:
                    cell = text(row_number, column, value, size=9, align="left" if column in (4, 5, 6) else "center")
                    cell.alignment = Alignment(horizontal=cell.alignment.horizontal, vertical="top", wrap_text=True)
                cell.fill = light if column < 10 else PatternFill("solid", fgColor="FFFFFF")
                cell.border = Border(left=thin if column == 1 else white, right=thin if column == 12 else white)
        last = start + len(statement.rows) - 1
        if last > start:
            sheet.merge_cells(start_row=start, start_column=1, end_row=last, end_column=1)
            sheet.cell(start, 1).alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
        # Group order and common receipt references just as in the reference workbook.
        group_start = start
        for offset in range(1, len(statement.rows) + 1):
            if offset == len(statement.rows) or statement.rows[offset].first_in_order:
                group_end = start + offset - 1
                if group_end > group_start:
                    for column, key in ((3, "sales_order"), (4, "delivery_receipt"), (5, "csi_invoice")):
                        values = {row[key] for row in statement.rows[group_start - start:offset]}
                        if len(values) == 1:
                            sheet.merge_cells(start_row=group_start, start_column=column, end_row=group_end, end_column=column)
                            sheet.cell(group_start, column).alignment = Alignment(horizontal="center" if column == 3 else "left", vertical="top", wrap_text=True)
                group_start = start + offset
        total_row = last + 1
        merged(total_row, 1, 9, "Grand Total", bold=True)
        sheet.row_dimensions[total_row].height = 23
        for column, key in ((10, "amount"), (11, "payment"), (12, "balance")):
            cell = sheet.cell(total_row, column, statement.totals[key])
            cell.font = Font(name="Arial", size=10, bold=True)
            cell.number_format = "#,##0.00"
            cell.alignment = Alignment(horizontal="right")
        for cell in sheet[total_row]:
            cell.fill = grey
            cell.border = Border(top=thin, bottom=thin)
        bank_row = total_row + 5
        bank_values = ["FOR ONLINE PAYMENT", f"{settings.get('bank_name') or ''} ACCOUNT: {settings.get('bank_account') or ''}", settings.get("bank_payee") or ""]
        for offset, value in enumerate(bank_values):
            row = bank_row + offset
            merged(row, 3, 11, value, size=12, bold=True, color="FF0000" if offset == 0 else "000000", align="center")
            sheet.row_dimensions[row].height = 22
            for column in range(3, 12):
                sheet.cell(row, column).border = Border(left=thin if column == 3 else Side(), right=thin if column == 11 else Side(),
                    top=thin if offset == 0 else Side(), bottom=thin if offset == 2 else Side())
        note_row = bank_row + 4
        notes = [settings.get("payment_reminder") or "", "Please make all checks payable to " + (settings.get("bank_payee") or ""),
                 "Please send us an email once payment has been made", "E-mail: " + (settings.get("contact_email") or "")]
        for offset, value in enumerate(notes):
            merged(note_row + offset, 1, 12, value, size=12, italic=True, color="FF0000" if offset == 1 else "000000")
            sheet.row_dimensions[note_row + offset].height = 22
        signature_row = note_row + 6
        for start_col, end_col, label, name, role in [
            (1, 4, "Prepared By:", settings.get("prepared_by"), settings.get("prepared_title")),
            (5, 8, "Check By:", settings.get("checked_by"), settings.get("checked_title")),
            (9, 12, "Conforme & Received By:", "____________________________", "Print Name and Signature / Date"),
        ]:
            merged(signature_row, start_col, end_col, label, size=9, bold=True)
            cell = merged(signature_row + 1, start_col, end_col, name or "____________________________", size=10, bold=True, align="center")
            cell.border = Border(bottom=thin)
            merged(signature_row + 2, start_col, end_col, role or "", size=9, align="center")
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_margins = PageMargins(left=.25, right=.25, top=.25, bottom=.25, header=0, footer=0)
        sheet.print_options.horizontalCentered = True
        sheet.print_title_rows = "1:9"
        sheet.print_area = f"A1:L{signature_row + 2}"
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
