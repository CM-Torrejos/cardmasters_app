import csv
import unittest
from datetime import date
from io import BytesIO, StringIO
from unittest.mock import patch

import frappe
from openpyxl import load_workbook

from cardmasters_app.cardmasters_app.services import soa_exports as exports


class TestSOAExports(unittest.TestCase):
    def setUp(self):
        self.previous_user = frappe.session.user
        frappe.set_user("Administrator")
        frappe.db.savepoint("soa_exports_test")

    def tearDown(self):
        frappe.db.rollback(save_point="soa_exports_test")
        frappe.set_user(self.previous_user)

    def test_csv_roundtrip_unicode_quotes_newlines_and_formula_text(self):
        columns = [("name", "Customer", "text"), ("amount", "Amount", "money"), ("date", "Date", "date")]
        groups = [{"rows": [['José, "Customer"\nSecond line', -25.5, date(2026, 10, 6)],
                             [" \t=1+1", 0, None], ["@SUM(A1)", 10, None]]}]
        content = exports.csv_content(columns, groups)
        self.assertTrue(content.startswith(b"\xef\xbb\xbf"))
        rows = list(csv.reader(StringIO(content.decode("utf-8-sig"))))
        self.assertEqual(rows[1], ['José, "Customer"\nSecond line', "-25.5", "2026-10-06"])
        self.assertEqual(rows[2][0], "' \t=1+1")
        self.assertEqual(rows[3][0], "'@SUM(A1)")

    def test_excel_types_sheet_names_and_formula_text(self):
        columns = [("reference", "Reference", "text"), ("amount", "Amount", "money"), ("date", "Date", "date")]
        groups = [{"customer_name": name, "rows": [["=1+1", -12.75, date(2026, 10, 6)], ["00012", 0, None]]}
                  for name in ["Client / Branch [1]" * 3, "Client / Branch [1]" * 3, ""]]
        workbook = load_workbook(BytesIO(exports.xlsx_content(columns, groups)))
        self.assertEqual(len(workbook.sheetnames), 3)
        self.assertEqual(len({name.casefold() for name in workbook.sheetnames}), 3)
        for sheet in workbook:
            self.assertLessEqual(len(sheet.title), 31)
            self.assertEqual(sheet["A2"].data_type, "s")
            self.assertEqual(sheet["A2"].value, "=1+1")
            self.assertEqual(sheet["A3"].value, "00012")
            self.assertEqual(sheet["B2"].value, -12.75)
            self.assertEqual(sheet["B2"].data_type, "n")
            self.assertEqual(sheet["C2"].value.date(), date(2026, 10, 6))
            self.assertEqual(sheet.freeze_panes, "A2")
            self.assertEqual(sheet.auto_filter.ref, "A1:C3")

    def test_rejects_invalid_format_before_loading_data(self):
        with patch.object(exports, "collect_export") as collect:
            with self.assertRaises(frappe.ValidationError):
                exports.download_export("statement", "html")
            collect.assert_not_called()

    def test_requires_statement_read_and_export_permission(self):
        for permission in ("read", "export"):
            fake_doc = frappe._dict()
            def check_permission(kind):
                if kind == permission:
                    raise frappe.PermissionError
            fake_doc.check_permission = check_permission
            with self.subTest(permission=permission), patch("frappe.get_doc", return_value=fake_doc):
                with self.assertRaises(frappe.PermissionError):
                    exports.collect_export("statement")

    def test_customer_selection_and_customer_permission_are_checked_before_reporting(self):
        doc = frappe.new_doc(exports.STATEMENT_DOCTYPE)
        doc.append("customers", {"customer": "allowed"})
        with patch("frappe.get_doc", return_value=doc), patch.object(doc, "check_permission"):
            with self.assertRaises(frappe.ValidationError):
                exports.collect_export("statement", customer="unrelated")
        customer = frappe._dict()
        customer.check_permission = lambda kind: (_ for _ in ()).throw(frappe.PermissionError)
        with patch("frappe.get_doc", side_effect=[doc, customer]), patch.object(doc, "check_permission"), patch.object(exports, "native") as native:
            with self.assertRaises(frappe.PermissionError):
                exports.collect_export("statement")
            native.assert_not_called()

    def test_real_native_reports_csv_and_excel_downloads(self):
        candidates = frappe.get_all("Sales Invoice", filters={"docstatus": 1, "outstanding_amount": [">", 0]},
            fields=["company", "customer", "posting_date"], order_by="posting_date desc", limit=1)
        if not candidates:
            self.skipTest("No submitted invoice with an outstanding balance is available")
        invoice = candidates[0]
        for report_type in exports.REPORT_TYPES:
            with self.subTest(report_type=report_type):
                doc = frappe.new_doc(exports.STATEMENT_DOCTYPE)
                doc.company = invoice.company
                doc.report = report_type
                doc.from_date = invoice.posting_date
                doc.to_date = invoice.posting_date
                doc.posting_date = invoice.posting_date
                doc.include_ageing = 1
                doc.append("customers", {"customer": invoice.customer,
                    "customer_name": frappe.db.get_value("Customer", invoice.customer, "customer_name")})
                doc.insert(set_name="Export test " + frappe.generate_hash(length=12))
                raw = exports.native().get_statement_dict(doc, get_statement_dict=True)
                loaded, columns, groups = exports.collect_export(doc.name)
                self.assertEqual(len(groups), 1)
                self.assertEqual(groups[0]["customer"], invoice.customer)
                rows, ageing = raw[invoice.customer]
                self.assertEqual(len(groups[0]["rows"]), len([row for row in rows if row]) + bool(ageing))
                keys = [key for key, label, kind in columns]
                actual = [dict(zip(keys, row, strict=True)) for row in groups[0]["rows"]]
                transactions = [row for row in actual if row["row_type"] == "Transaction"]
                source = [row for row in rows if row.get("voucher_no")]
                amount_key = "debit" if report_type == "General Ledger" else "outstanding"
                self.assertEqual([row[amount_key] for row in transactions], [row.get(amount_key) for row in source])
                for row in actual:
                    self.assertEqual(row["customer"], invoice.customer)
                    self.assertEqual(row["to_date"], invoice.posting_date)
                filtered = exports.collect_export(doc.name, customer=invoice.customer)[2]
                self.assertEqual(filtered, groups)
                for file_format in ("csv", "xlsx"):
                    exports.download_export(doc.name, file_format)
                    content = frappe.local.response.filecontent
                    self.assertEqual(frappe.local.response.type, "download")
                    self.assertTrue(frappe.local.response.filename.endswith("." + file_format))
                    if file_format == "csv":
                        parsed = list(csv.reader(StringIO(content.decode("utf-8-sig"))))
                        self.assertEqual(len(parsed), len(actual) + 1)
                    else:
                        workbook = load_workbook(BytesIO(content))
                        self.assertEqual(workbook.active.max_row, len(actual) + 1)
                        amount_cell = workbook.active.cell(2, keys.index(amount_key) + 1)
                        self.assertEqual(amount_cell.data_type, "n")

    def test_no_transactions_returns_clear_error(self):
        doc = frappe.new_doc(exports.STATEMENT_DOCTYPE)
        doc.report = "General Ledger"
        doc.append("customers", {"customer": "allowed"})
        customer = frappe._dict(customer_name="Allowed", check_permission=lambda kind: None)
        with patch("frappe.get_doc", side_effect=[doc, customer]), patch.object(doc, "check_permission"), patch.object(exports, "native") as native:
            native.return_value.get_statement_dict.return_value = {}
            with self.assertRaisesRegex(frappe.ValidationError, "No statement transactions"):
                exports.collect_export("statement")

    def test_single_customer_filter_is_passed_to_native_report(self):
        doc = frappe.new_doc(exports.STATEMENT_DOCTYPE)
        doc.report = "Accounts Receivable"
        doc.company = "Company"
        doc.include_ageing = 0
        for name in ("first", "second"):
            doc.append("customers", {"customer": name})
        customer = frappe._dict(customer_name="Second", check_permission=lambda kind: None)
        with patch("frappe.get_doc", side_effect=[doc, customer]), patch.object(doc, "check_permission"), patch.object(exports, "native") as native:
            native.return_value.get_statement_dict.return_value = {
                "second": [[{"voucher_no": "PE-1", "voucher_type": "Payment Entry", "currency": "PHP", "outstanding": -10}], []],
            }
            groups = exports.collect_export("statement", customer="second")[2]
            reported_doc = native.return_value.get_statement_dict.call_args.args[0]
            self.assertEqual([row.customer for row in reported_doc.customers], ["second"])
            self.assertEqual([row.customer for row in doc.customers], ["first", "second"])
            self.assertEqual(groups[0]["customer"], "second")
