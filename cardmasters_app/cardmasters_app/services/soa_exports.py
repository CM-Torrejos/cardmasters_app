"""Structured CSV/XLSX statements from the native SOA report data."""

import copy
import csv
import re
from datetime import date, datetime
from io import BytesIO, StringIO

import frappe
from frappe import _
from frappe.utils import getdate

from cardmasters_app.cardmasters_app.services.soa_templates import REPORT_TYPES, STATEMENT_DOCTYPE, native

COMMON_COLUMNS = [
    ("company", "Company", "text"),
    ("customer", "Customer ID", "text"),
    ("customer_name", "Customer Name", "text"),
    ("currency", "Currency", "text"),
    ("from_date", "From Date", "date"),
    ("to_date", "To Date / As Of", "date"),
    ("row_type", "Row Type", "text"),
    ("description", "Description", "text"),
    ("posting_date", "Posting Date", "date"),
    ("voucher_type", "Voucher Type", "text"),
    ("voucher_no", "Reference", "text"),
    ("po_no", "Customer PO", "text"),
    ("due_date", "Due Date", "date"),
    ("remarks", "Remarks", "text"),
]
AGEING_COLUMNS = [
    ("range1", "0–30 Days", "money"), ("range2", "31–60 Days", "money"),
    ("range3", "61–90 Days", "money"), ("range4", "91–120 Days", "money"),
    ("range5", "Above 120 Days", "money"),
]


def export_columns(doc):
    amounts = [("debit", "Debit", "money"), ("credit", "Credit", "money"),
               ("balance", "Balance", "money")] if doc.report == "General Ledger" else [
        ("invoiced", "Invoiced", "money"), ("paid", "Paid", "money"),
        ("credit_note", "Credit Notes", "money"), ("outstanding", "Outstanding", "money"),
        ("age", "Age (Days)", "number"),
    ]
    return COMMON_COLUMNS + amounts + (AGEING_COLUMNS if doc.include_ageing else [])


def collect_export(document_name, customer=None):
    from erpnext import get_company_currency
    from erpnext.accounts.party import get_party_account_currency

    doc = frappe.get_doc(STATEMENT_DOCTYPE, document_name)
    doc.check_permission("read")
    doc.check_permission("export")
    if doc.report not in REPORT_TYPES:
        frappe.throw(_("Select a supported statement report."))
    if not doc.customers:
        frappe.throw(_("Add customers to the saved statement first."))
    selected = [row for row in doc.customers if not customer or row.customer == customer]
    if not selected:
        frappe.throw(_("Select a customer included in this statement."))
    customers = {}
    for entry in selected:
        customer_doc = frappe.get_doc("Customer", entry.customer)
        customer_doc.check_permission("read")
        customers[entry.customer] = customer_doc
    from cardmasters_app.cardmasters_app.services.soa_templates import get_selected_template
    template = get_selected_template(doc, doc.report, check_permission=True)
    if template and template.get("data_source") == "Sales Orders":
        from cardmasters_app.cardmasters_app.services.soa_sales_orders import collect_statement
        columns = [("customer_name", "Customer", "text"), ("invoice_date", "Invoice Date", "date"),
                   ("sales_order", "Sales Order", "text"), ("delivery_receipt", "Delivery Receipt", "text"),
                   ("csi_invoice", "CSI Invoice", "text"), ("item_description", "Item Description", "text"),
                   ("po_no", "PO", "text"), ("aged", "Aged", "text"), ("remarks", "Remarks", "text"),
                   ("amount", "Amount", "money"), ("payment", "Payment", "money"), ("balance", "Balance", "money"),
                   ("currency", "Currency", "text"), ("row_type", "Row Type", "text")]
        groups = []
        for customer_id in customers:
            statement = collect_statement(doc, customer_id, template)
            if statement:
                rows = [dict(row, currency=statement.currency, row_type="Transaction") for row in statement.rows]
                rows.append(dict(statement.totals, customer_name=statement.customer.customer_name,
                                 currency=statement.currency, row_type="Summary", item_description="Grand Total"))
                groups.append({"customer": customer_id, "customer_name": statement.customer.customer_name,
                    "rows": [[cell_value(row.get(key), kind) for key, label, kind in columns] for row in rows],
                    "statement": statement})
        if not groups:
            frappe.throw(_("No outstanding Sales Orders were found for these customers and statement date."))
        return doc, columns, groups
    report_doc = copy.deepcopy(doc)
    report_doc.set("customers", [row.as_dict() for row in selected])
    # This native mode returns the very same rows used for PDF generation, without
    # rendering HTML, loading the selected layout, or calculating balances again.
    statements = native().get_statement_dict(report_doc, get_statement_dict=True)
    if not statements:
        frappe.throw(_("No statement transactions were found for these customers and filters."))
    columns = export_columns(doc)
    groups = []
    for customer_id, (data, ageing) in statements.items():
        customer_doc = customers[customer_id]
        currency = (doc.currency or get_party_account_currency("Customer", customer_id, doc.company)
                    or get_company_currency(doc.company)) if doc.report == "General Ledger" else (
            next((row.get("currency") for row in data if row.get("currency")), None)
            or get_company_currency(doc.company)
        )
        invoice_names = list({row.get("voucher_no") for row in data
                              if row.get("voucher_type") == "Sales Invoice" and row.get("voucher_no")})
        invoices = {}
        if invoice_names:
            invoices = {row.name: row for row in frappe.get_all("Sales Invoice", filters={
                "name": ["in", invoice_names], "company": doc.company, "customer": customer_id,
            }, fields=["name", "po_no", "due_date"])}
        metadata = {
            "company": doc.company, "customer": customer_id, "customer_name": customer_doc.customer_name,
            "currency": currency,
            "from_date": doc.from_date if doc.report == "General Ledger" else None,
            "to_date": doc.to_date if doc.report == "General Ledger" else doc.posting_date,
        }
        rows = []
        for source in data:
            if not source:
                continue
            row = dict(source)
            invoice = invoices.get(row.get("voucher_no"), {})
            row.update(metadata)
            row["row_type"] = "Transaction" if source.get("voucher_no") else "Summary"
            row["description"] = "" if source.get("voucher_no") else (
                source.get("account") or source.get("party") or _("Total"))
            row["po_no"] = invoice.get("po_no")
            # A payment-term due date from the native report takes precedence.
            row["due_date"] = source.get("due_date") or invoice.get("due_date")
            rows.append([cell_value(row.get(key), kind) for key, label, kind in columns])
        if doc.include_ageing and ageing:
            row = {key: ageing[0].get(key) for key, label, kind in AGEING_COLUMNS}
            row.update(metadata)
            # Native ageing is calculated independently of the GL presentation currency.
            row["currency"] = ageing[0].get("currency") or get_company_currency(doc.company)
            row["row_type"] = "Ageing Summary"
            row["description"] = _("Ageing based on {0}").format(doc.ageing_based_on)
            rows.append([cell_value(row.get(key), kind) for key, label, kind in columns])
        groups.append({"customer": customer_id, "customer_name": customer_doc.customer_name, "rows": rows})
    return doc, columns, groups


def cell_value(value, kind):
    if value is None or value == "":
        return None
    if kind == "date":
        return getdate(value)
    # Report numbers stay numeric; textual identifiers retain leading zeroes.
    return str(value) if kind == "text" else value


def csv_value(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, str) and value.lstrip(" \t\r\n").startswith(("=", "+", "-", "@")):
        # A text field must not become a formula when opened in spreadsheet software.
        return "'" + value
    return value


def csv_content(columns, groups):
    stream = StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow([label for key, label, kind in columns])
    for group in groups:
        writer.writerows([[csv_value(value) for value in row] for row in group["rows"]])
    return stream.getvalue().encode("utf-8-sig")


def sheet_title(name, used):
    name = re.sub(r"[\x00-\x1f\\/*?:\[\]]", " ", name).strip(" '") or "Customer"
    base = name[:31]
    title = base
    suffix = 2
    while title.casefold() in used:
        ending = f" ({suffix})"
        title = base[:31 - len(ending)] + ending
        suffix += 1
    used.add(title.casefold())
    return title


def xlsx_content(columns, groups):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    from frappe.utils.xlsxutils import ILLEGAL_CHARACTERS_RE

    workbook = Workbook(write_only=True)
    used = set()
    for group in groups:
        sheet = workbook.create_sheet(sheet_title(group["customer_name"], used))
        sheet.freeze_panes = "A2"
        for index, (key, label, kind) in enumerate(columns, 1):
            sheet.column_dimensions[get_column_letter(index)].width = (
                38 if key in ("remarks", "description") else 26 if key in ("customer_name", "voucher_no") else 18)
        for row_index, values in enumerate([[label for key, label, kind in columns], *group["rows"]]):
            cells = []
            for value, (key, label, kind) in zip(values, columns, strict=True):
                if isinstance(value, str):
                    value = ILLEGAL_CHARACTERS_RE.sub("", value)
                cell = WriteOnlyCell(sheet, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"
                if row_index == 0:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="15375B")
                elif kind == "date":
                    cell.number_format = "yyyy-mm-dd"
                elif kind == "money":
                    cell.number_format = "#,##0.00;[Red]-#,##0.00"
                cells.append(cell)
            sheet.append(cells)
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(group['rows']) + 1}"
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


@frappe.whitelist()
def download_export(document_name, file_format="xlsx", customer=None, layout="template"):
    file_format = (file_format or "").lower()
    if file_format not in ("csv", "xlsx"):
        frappe.throw(_("Choose CSV or Excel (.xlsx)."))
    if layout not in ("template", "data"):
        frappe.throw(_("Choose the template layout or data table layout."))
    doc, columns, groups = collect_export(document_name, customer=customer)
    from cardmasters_app.cardmasters_app.services.soa_templates import get_selected_template
    template = get_selected_template(doc, doc.report, check_permission=True)
    if file_format == "xlsx" and layout == "template" and template and template.get("excel_layout") == "Sales Order Statement":
        from cardmasters_app.cardmasters_app.services.soa_excel_layouts import designed_workbook
        content = designed_workbook([group["statement"] for group in groups])
    else:
        content = csv_content(columns, groups) if file_format == "csv" else xlsx_content(columns, groups)
    filename = re.sub(r'[\x00-\x1f\\/:*?"<>|]', "_", doc.name).strip(" .") or "Statement of Accounts"
    frappe.local.response.filename = f"{filename}.{file_format}"
    frappe.local.response.filecontent = content
    frappe.local.response.content_type = ("text/csv; charset=utf-8" if file_format == "csv" else
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    frappe.local.response.type = "download"
