import unittest
from decimal import Decimal
from datetime import date
from io import BytesIO
from unittest.mock import patch

import frappe
from openpyxl import load_workbook

from cardmasters_app.cardmasters_app.services import soa_sales_orders as sales
from cardmasters_app.cardmasters_app.services.soa_excel_layouts import designed_workbook
from cardmasters_app.cardmasters_app.services import soa_templates as templates


def row(**values):
    return frappe._dict(values)


class TestSalesOrderSOA(unittest.TestCase):
    def setUp(self):
        self.previous_user = frappe.session.user
        frappe.set_user("Administrator")
        frappe.db.savepoint("sales_order_soa_test")

    def tearDown(self):
        frappe.db.rollback(save_point="sales_order_soa_test")
        frappe.set_user(self.previous_user)

    def test_proportional_amounts_and_payments_preserve_every_cent(self):
        amounts = sales.allocate(100, [60, 40])
        payments = sales.allocate(25, amounts)
        self.assertEqual(amounts, [Decimal("60.00"), Decimal("40.00")])
        self.assertEqual(payments, [Decimal("15.00"), Decimal("10.00")])
        for total in ("0.04", "0.05", "0.01"):
            allocated = sales.allocate(total, [1] * 6)
            self.assertEqual(sum(allocated), Decimal(total))
            self.assertTrue(all(value >= 0 for value in allocated))
        self.assertEqual(sum(sales.allocate(-1, [2, 3])), Decimal("-1.00"))

    def test_shared_invoice_and_unlinked_items_are_not_double_counted(self):
        invoice = row(name="SI", posted_total=100, is_return=0)
        items = [row(parent="SI", sales_order="SO1", base_net_amount=60),
                 row(parent="SI", sales_order="SO2", base_net_amount=30),
                 row(parent="SI", sales_order=None, base_net_amount=10)]
        ledger = [row(against_voucher_no="SI", voucher_type="Sales Invoice", voucher_no="SI", amount=100),
                  row(against_voucher_no="SI", voucher_type="Payment Entry", voucher_no="PE", amount=-50)]
        self.assertEqual(sales.paid_as_of({"SO1", "SO2"}, [invoice], items, ledger, [], {}),
                         {"SO1": Decimal("30.00"), "SO2": Decimal("15.00")})

    def test_transferred_advance_is_counted_once(self):
        invoice = row(name="SI", posted_total=100, is_return=0)
        items = [row(parent="SI", sales_order="SO", base_net_amount=100)]
        ledger = [row(against_voucher_no="SI", voucher_type="Sales Invoice", voucher_no="SI", amount=100),
                  row(against_voucher_no="SI", voucher_type="Payment Entry", voucher_no="PE", amount=-30)]
        direct = [row(sales_order="SO", voucher_type="Payment Entry", voucher_no="PE", amount=30)]
        advances = {"SI": [row(reference_type="Payment Entry", reference_name="PE", allocated_amount=30)]}
        self.assertEqual(sales.paid_as_of({"SO"}, [invoice], items, ledger, direct, advances)["SO"], Decimal("30.00"))
        # Independent invoice and SO allocations from the same payment are both counted.
        self.assertEqual(sales.paid_as_of({"SO"}, [invoice], items, ledger, direct, {})["SO"], Decimal("60.00"))

    def test_payment_before_a_future_invoice_is_an_advance_at_cutoff(self):
        invoice = row(name="SI", posted_total=0, is_return=0)
        items = [row(parent="SI", sales_order="SO", base_net_amount=100)]
        ledger = [row(against_voucher_no="SI", voucher_type="Payment Entry", voucher_no="PE", amount=-20)]
        self.assertEqual(sales.paid_as_of({"SO"}, [invoice], items, ledger, [], {})["SO"], Decimal("20.00"))

    def test_applied_and_unapplied_returns_reduce_balance_once(self):
        invoices = [row(name="SI", posted_total=100, is_return=0), row(name="RETURN", posted_total=-10, is_return=1)]
        items = [row(parent="SI", sales_order="SO", base_net_amount=100),
                 row(parent="RETURN", sales_order="SO", base_net_amount=-10)]
        unapplied = [row(against_voucher_no="SI", voucher_type="Sales Invoice", voucher_no="SI", amount=100),
                     row(against_voucher_no="RETURN", voucher_type="Sales Invoice", voucher_no="RETURN", amount=-10)]
        applied = unapplied + [row(against_voucher_no="SI", voucher_type="Sales Invoice", voucher_no="RETURN", amount=-10),
                               row(against_voucher_no="RETURN", voucher_type="Sales Invoice", voucher_no="RETURN", amount=10)]
        for ledger in (unapplied, applied):
            self.assertEqual(sales.paid_as_of({"SO"}, invoices, items, ledger, [], {})["SO"], Decimal("10.00"))

    def test_due_date_follows_first_unpaid_payment_term(self):
        terms = [row(due_date=date(2026, 1, 1), payment_amount=30), row(due_date=date(2026, 2, 1), payment_amount=70)]
        self.assertEqual(sales.item_due_date(terms, Decimal(30), []), date(2026, 2, 1))
        self.assertIsNone(sales.item_due_date([], Decimal(0), []))

    def sample_statement(self):
        template = frappe.get_doc("SOA Template", "CardMasters Sales Order SOA")
        return row(customer=row(customer_name="Sample Customer"), company=row(company_name="CardMasters"),
            date=date(2026, 10, 6), number="46300", currency="PHP", settings=template.as_dict(),
            logo_url="", rows=[row(customer_name="Sample Customer", sales_order="SO-1", invoice_date=date(2026, 3, 11),
                delivery_receipt="23384", csi_invoice="8303", item_description="2 PCS VINYL ON SINTRA", po_no="PO-2026-04",
                aged="Past Due", remarks="PARTLY PAID", amount=60, payment=15, balance=45, first_in_order=True),
                row(customer_name="Sample Customer", sales_order="SO-1", invoice_date=date(2026, 3, 11),
                delivery_receipt="23384", csi_invoice="8303", item_description="2 PCS VINYL STICKER", po_no="PO-2026-04",
                aged="Past Due", remarks="PARTLY PAID", amount=40, payment=10, balance=30, first_in_order=False)],
            totals=row(amount=100, payment=25, balance=75))

    def test_designed_excel_has_layout_numeric_cells_logo_and_print_settings(self):
        statement = self.sample_statement()
        workbook = load_workbook(BytesIO(designed_workbook([statement])))
        sheet = workbook.active
        self.assertEqual(sheet["D2"].value, "STATEMENT OF ACCOUNT")
        self.assertEqual(sheet["J9"].fill.fgColor.rgb, "008396B5")
        self.assertEqual(sheet["J10"].value, 60)
        self.assertEqual(sheet["K10"].value, 15)
        self.assertEqual(sheet["L10"].value, 45)
        self.assertEqual(sheet["L12"].value, 75)
        self.assertEqual(sheet["J10"].data_type, "n")
        self.assertIn("A10:A11", str(sheet.merged_cells))
        self.assertIn("C10:C11", str(sheet.merged_cells))
        self.assertEqual(sheet.page_setup.orientation, "landscape")
        self.assertEqual(sheet.page_setup.fitToWidth, 1)
        self.assertEqual(sheet.print_title_rows, "$1:$9")
        self.assertEqual(len(sheet._images), 1)
        self.assertFalse(sheet.sheet_view.showGridLines)

    def test_layout_text_stays_text_in_excel(self):
        statement = self.sample_statement()
        statement.customer.customer_name = "=1+1"
        statement.rows[0].po_no = "000123"
        statement.rows[0].item_description = "=1+1"
        sheet = load_workbook(BytesIO(designed_workbook([statement]))).active
        self.assertEqual(sheet["C7"].data_type, "s")
        self.assertEqual(sheet["F10"].data_type, "s")
        self.assertEqual(sheet["G10"].value, "000123")

    def test_source_and_excel_layout_combinations_are_validated(self):
        template = frappe.get_doc("SOA Template", "CardMasters Sales Order SOA")
        template.report_type = "General Ledger"
        with self.assertRaises(frappe.ValidationError):
            templates.validate_template_document(template)

    def test_real_unbilled_order_appears_without_invoice_dependency(self):
        # Local financial records remain untouched; only a rollback-only statement is saved.
        orders = frappe.get_all("Sales Order", filters={"docstatus": 1, "per_billed": 0},
            fields=["name", "customer", "company", "transaction_date"], order_by="transaction_date desc", limit=1)
        if not orders:
            self.skipTest("No submitted unbilled order available")
        order = orders[0]
        doc = frappe.new_doc("Process Statement Of Accounts")
        doc.company = order.company
        doc.report = "Accounts Receivable"
        doc.posting_date = order.transaction_date
        doc.custom_soa_template = "CardMasters Sales Order SOA"
        doc.append("customers", {"customer": order.customer,
            "customer_name": frappe.db.get_value("Customer", order.customer, "customer_name"),
            "billing_email": "soa-test@example.invalid"})
        doc.insert(set_name="Sales Order SOA test " + frappe.generate_hash(length=12))
        template = frappe.get_doc("SOA Template", doc.custom_soa_template)
        statement = sales.collect_statement(doc, order.customer, template)
        self.assertIsNotNone(statement)
        self.assertTrue(any(line.sales_order == order.name for line in statement.rows))
        self.assertEqual(sum(sales.money(line.balance) for line in statement.rows), sales.money(statement.totals.balance))
        html = templates.preview(doc.name)
        self.assertIn("STATEMENT OF ACCOUNT", html)
        self.assertIn("data:image/png;base64,", html)
        from cardmasters_app.cardmasters_app.services.soa_exports import download_export
        download_export(doc.name)
        workbook = load_workbook(BytesIO(frappe.local.response.filecontent))
        self.assertEqual(workbook.active["D2"].value, "STATEMENT OF ACCOUNT")
        from cardmasters_app.cardmasters_app.services.soa_sales_order_delivery import report_pdf
        self.assertTrue(report_pdf(doc, template).startswith(b"%PDF"))
        with patch("frappe.enqueue") as enqueue:
            self.assertTrue(templates.send_emails(doc.name))
        self.assertTrue(enqueue.called)
        self.assertTrue(enqueue.call_args.kwargs["attachments"][0]["fcontent"].startswith(b"%PDF"))

    def test_shared_invoice_rounding_preserves_small_payments_and_free_items(self):
        invoice = row(name="SI", posted_total=3, is_return=0)
        items = [row(parent="SI", sales_order=name, base_net_amount=1) for name in ("SO1", "SO2", "SO3")]
        items.append(row(parent="SI", sales_order="FREE", base_net_amount=0, base_amount=100))
        ledger = [row(against_voucher_no="SI", voucher_type="Sales Invoice", voucher_no="SI", amount=3),
                  row(against_voucher_no="SI", voucher_type="Payment Entry", voucher_no="PE", amount=-.01)]
        paid = sales.paid_as_of({"SO1", "SO2", "SO3", "FREE"}, [invoice], items, ledger, [], {})
        self.assertEqual(sum(paid.values()), Decimal("0.01"))
        self.assertEqual(paid["FREE"], Decimal("0.00"))

    def test_cutoff_is_applied_to_payment_ledger_and_direct_payment_queries(self):
        cutoff = date(2026, 3, 1)
        invoice = row(name="FUTURE", posting_date=date(2026, 4, 1), is_return=0,
                      base_grand_total=100, base_rounded_total=100, disable_rounded_total=0)
        invoice_doc = row(check_permission=lambda kind: None)
        with patch("frappe.get_all", side_effect=[
            [row(parent="FUTURE", sales_order="SO", so_detail="ITEM")], [invoice],
            [row(parent="FUTURE", sales_order="SO", so_detail="ITEM", base_net_amount=100, base_amount=100)], [], [],
        ]) as get_all, patch("frappe.get_doc", return_value=invoice_doc), patch("frappe.db.sql", return_value=[]) as sql:
            result = sales.linked_data({"SO"}, "Company", "Customer", cutoff)
        self.assertEqual(result[0][0].posted_total, 0)
        invoice_call = next(call for call in get_all.call_args_list if call.args[0] == "Sales Invoice")
        self.assertIn("custom_dr_billing_reference", invoice_call.kwargs["fields"])
        self.assertIn("custom_bir_series", invoice_call.kwargs["fields"])
        ledger_call = next(call for call in get_all.call_args_list if call.args[0] == "Payment Ledger Entry")
        self.assertEqual(ledger_call.kwargs["filters"]["posting_date"], ["<=", cutoff])
        self.assertEqual(sql.call_args.args[1][2], cutoff)
        self.assertEqual(sql.call_args.args[1][6], cutoff)

    def test_invoice_reference_variant_preserves_original_mapping_and_item_cutoff(self):
        cutoff = date(2026, 3, 1)
        doc = row(company="Company", posting_date=cutoff, cost_center=[], name="STATEMENT")
        customer = row(customer_name="Customer", check_permission=lambda kind: None)
        customer.as_dict = lambda: row(customer_name="Customer")
        company = row(check_permission=lambda kind: None)
        company.as_dict = lambda: row(company_name="Company")
        order = row(name="SO", transaction_date=cutoff, currency="PHP", grand_total=100,
                    rounded_total=100, disable_rounded_total=0)
        items = [row(name=name, parent="SO", item_name=name, qty=1, uom="PCS", net_amount=50)
                 for name in ("ITEM1", "ITEM2")]
        invoices = [
            row(name="SI1", posting_date=cutoff, is_return=0, custom_bir_series="000123",
                custom_dr_billing_reference="000456"),
            row(name="SI2", posting_date=cutoff, is_return=0, custom_bir_series="000123",
                custom_dr_billing_reference="000456"),
            row(name="BLANK", posting_date=cutoff, is_return=0, custom_bir_series="",
                custom_dr_billing_reference=""),
            row(name="FUTURE", posting_date=date(2026, 3, 2), is_return=0,
                custom_bir_series="FUTURE-BIR", custom_dr_billing_reference="FUTURE-DR"),
            row(name="RETURN", posting_date=cutoff, is_return=1,
                custom_bir_series="RETURN-BIR", custom_dr_billing_reference="RETURN-DR"),
        ]
        links = [row(parent=invoice.name, sales_order="SO", so_detail="ITEM1") for invoice in invoices]
        delivery_items = [row(parent="DN", so_detail="ITEM1")]
        with patch("frappe.get_doc", side_effect=[customer, company]), \
                patch("erpnext.get_company_currency", return_value="PHP"), \
                patch("frappe.get_list", return_value=[order]), \
                patch("frappe.get_all", side_effect=[items, [], [("SO", "")], delivery_items,
                    [row(name="DN", custom_reference_no="ORIGINAL-DR")]]), \
                patch.object(sales, "linked_data", return_value=(invoices, links, [], [], {})), \
                patch.object(sales, "paid_as_of", return_value={"SO": Decimal(0)}):
            statement = sales.collect_statement(doc, "Customer", row(as_dict=lambda: {}))
        first, second = statement.rows
        self.assertEqual(first.dr_billing_reference, "000456")
        self.assertEqual(first.bir_series, "000123")
        self.assertEqual(first.delivery_receipt, "ORIGINAL-DR")
        self.assertEqual(first.csi_invoice, "BLANK, 000123")
        self.assertEqual(second.dr_billing_reference, "")
        self.assertEqual(second.bir_series, "")
        self.assertEqual(statement.totals.balance, 100)
        from pathlib import Path
        html = Path(frappe.get_app_path("cardmasters_app", "templates", "soa",
                                        "sales_order_invoice_references.html")).read_text()
        rendered = templates.template_environment().from_string(html).render(statement=statement)
        self.assertIn("000456", rendered)
        self.assertIn("000123", rendered)
        for excluded in ("ORIGINAL-DR", "BLANK", "FUTURE-DR", "RETURN-DR"):
            self.assertNotIn(excluded, rendered)
